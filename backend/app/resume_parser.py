import io
from typing import Tuple
try:
    import pymupdf as fitz
except ImportError:
    import fitz  # PyMuPDF fallback
import docx


MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB


class ResumeParseError(Exception):
    """Custom exception for resume parsing issues with clear error message."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def validate_file_metadata(filename: str, file_size: int) -> str:
    """
    Validates file extension and size.
    Returns detected extension ('pdf' or 'docx').
    """
    if file_size > MAX_FILE_SIZE_BYTES:
        raise ResumeParseError(
            f"File size exceeds the 5 MB limit (received {file_size / (1024 * 1024):.2f} MB). "
            "Please upload a smaller document.",
            status_code=413
        )
    
    if file_size == 0:
        raise ResumeParseError("The uploaded file is empty (0 bytes).", status_code=400)

    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        return "pdf"
    elif filename_lower.endswith(".docx"):
        return "docx"
    elif filename_lower.endswith(".doc"):
        raise ResumeParseError(
            "Legacy .doc format is not directly supported. Please save your file as .docx or .pdf and re-upload.",
            status_code=400
        )
    else:
        raise ResumeParseError(
            "Unsupported file format. Please upload a PDF or DOCX resume document.",
            status_code=400
        )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using PyMuPDF (fitz).
    Checks for corrupt files and image-only/scanned documents without text.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ResumeParseError(
            f"Unable to read PDF file. The file may be corrupted or password-protected. ({str(e)})",
            status_code=400
        )

    if doc.page_count == 0:
        raise ResumeParseError("The PDF document contains no pages.", status_code=400)

    extracted_pages = []
    for page_num in range(doc.page_count):
        page = doc.load_page(page_num)
        text = page.get_text("text")
        if text.strip():
            extracted_pages.append(text.strip())

    combined_text = "\n\n".join(extracted_pages).strip()

    if not combined_text:
        raise ResumeParseError(
            "No extractable text was found in this PDF. The document appears to be a scanned image "
            "or raster graphics. Optical Character Recognition (OCR) is not yet supported in this phase. "
            "Please upload a text-based PDF or DOCX file.",
            status_code=422
        )

    return combined_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX bytes using python-docx.
    Traverses paragraphs and tables.
    """
    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
    except Exception as e:
        raise ResumeParseError(
            f"Unable to read DOCX file. The file may be corrupted or invalid. ({str(e)})",
            status_code=400
        )

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Also capture text within tables if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in paragraphs:
                    paragraphs.append(cell_text)

    combined_text = "\n".join(paragraphs).strip()

    if not combined_text:
        raise ResumeParseError(
            "No extractable text was found in this DOCX document. Please check that the file contains text content.",
            status_code=422
        )

    return combined_text


def parse_resume_bytes(filename: str, file_bytes: bytes) -> Tuple[str, str, int]:
    """
    Validates and extracts raw text from resume bytes.
    Returns: (extracted_text, file_type, word_count)
    """
    file_type = validate_file_metadata(filename, len(file_bytes))

    if file_type == "pdf":
        text = extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        text = extract_text_from_docx(file_bytes)
    else:
        raise ResumeParseError("Unsupported file type.", status_code=400)

    word_count = len(text.split())
    return text, file_type, word_count
