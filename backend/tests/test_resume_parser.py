"""
HerTurn Phase 4: Resume Parsing Security & Boundary Validation Test Suite.
Verifies:
- PDF and DOCX text extraction
- Multi-page PDF handling and text aggregation
- DOCX paragraph and table parsing
- File size boundary enforcement (5MB limit -> 413)
- Empty file rejection (0 bytes -> 400)
- Unsupported format rejection (.txt, .doc, .png, .exe -> 400)
- Corrupted document detection (corrupted PDF/DOCX -> 400)
- Scanned / image-only PDF detection (no extractable text -> 422)
- Fast-path API endpoint /api/resume/parse validation
"""

import io
import pytest
import pymupdf as fitz
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.resume_parser import (
    validate_file_metadata,
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_resume_bytes,
    ResumeParseError,
    MAX_FILE_SIZE_BYTES,
)

client = TestClient(app)


def create_sample_pdf(text: str) -> bytes:
    """Helper to generate a valid in-memory PDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(paragraphs: list[str], table_cells: list[str] = None) -> bytes:
    """Helper to generate a valid in-memory DOCX with paragraphs and optional tables."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_cells:
        table = doc.add_table(rows=1, cols=len(table_cells))
        for idx, cell_text in enumerate(table_cells):
            table.cell(0, idx).text = cell_text
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# -----------------------------------------------------------------------------
# Metadata & Boundary Validation Tests
# -----------------------------------------------------------------------------

def test_validate_file_metadata_pdf_success():
    ext = validate_file_metadata("Candidate_Resume.pdf", 1024)
    assert ext == "pdf"


def test_validate_file_metadata_docx_success():
    ext = validate_file_metadata("Candidate_Resume.docx", 2048)
    assert ext == "docx"


def test_validate_file_metadata_zero_bytes_rejected():
    with pytest.raises(ResumeParseError) as exc_info:
        validate_file_metadata("empty.pdf", 0)
    assert exc_info.value.status_code == 400
    assert "empty" in exc_info.value.message.lower()


def test_validate_file_metadata_size_exceeded_rejected():
    oversized = MAX_FILE_SIZE_BYTES + 1024
    with pytest.raises(ResumeParseError) as exc_info:
        validate_file_metadata("large.pdf", oversized)
    assert exc_info.value.status_code == 413
    assert "5 mb limit" in exc_info.value.message.lower()


def test_validate_file_metadata_legacy_doc_rejected():
    with pytest.raises(ResumeParseError) as exc_info:
        validate_file_metadata("legacy_cv.doc", 5000)
    assert exc_info.value.status_code == 400
    assert "legacy .doc" in exc_info.value.message.lower()


def test_validate_file_metadata_unsupported_extensions_rejected():
    for filename in ["resume.txt", "cv.png", "malware.exe", "document.rtf"]:
        with pytest.raises(ResumeParseError) as exc_info:
            validate_file_metadata(filename, 5000)
        assert exc_info.value.status_code == 400
        assert "unsupported file format" in exc_info.value.message.lower()


# -----------------------------------------------------------------------------
# PDF Parsing & Corrupted/Scanned Detection
# -----------------------------------------------------------------------------

def test_extract_text_from_valid_pdf():
    sample_text = "Anita Sharma - Senior Cloud Architect\n10 years experience in distributed systems."
    pdf_bytes = create_sample_pdf(sample_text)
    extracted = extract_text_from_pdf(pdf_bytes)
    assert "Anita Sharma" in extracted
    assert "Senior Cloud Architect" in extracted


def test_extract_text_from_multipage_pdf():
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Page 1: Work Experience at Infosys")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Page 2: Education and Cloud Certifications")
    pdf_bytes = doc.tobytes()
    doc.close()

    extracted = extract_text_from_pdf(pdf_bytes)
    assert "Page 1" in extracted
    assert "Page 2" in extracted


def test_extract_text_from_corrupted_pdf_rejected():
    corrupted_bytes = b"%PDF-1.4\nCorrupted binary junk \x00\xff\xfe random non-pdf stream"
    with pytest.raises(ResumeParseError) as exc_info:
        extract_text_from_pdf(corrupted_bytes)
    assert exc_info.value.status_code == 400
    assert "unable to read pdf" in exc_info.value.message.lower()


def test_extract_text_from_scanned_or_blank_pdf_rejected():
    # Empty page with no text
    doc = fitz.open()
    doc.new_page()  # Blank page
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(ResumeParseError) as exc_info:
        extract_text_from_pdf(pdf_bytes)
    assert exc_info.value.status_code == 422
    assert "scanned image" in exc_info.value.message.lower()


# -----------------------------------------------------------------------------
# DOCX Parsing & Edge Cases
# -----------------------------------------------------------------------------

def test_extract_text_from_valid_docx():
    paragraphs = [
        "Priya Raman - Full Stack Developer",
        "Experience: 5 years building scalable web services in React and Python.",
    ]
    docx_bytes = create_sample_docx(paragraphs)
    extracted = extract_text_from_docx(docx_bytes)
    assert "Priya Raman" in extracted
    assert "Full Stack Developer" in extracted


def test_extract_text_from_docx_with_tables():
    paragraphs = ["Experience Summary:"]
    table_skills = ["Python", "FastAPI", "PostgreSQL", "Docker"]
    docx_bytes = create_sample_docx(paragraphs, table_cells=table_skills)
    extracted = extract_text_from_docx(docx_bytes)
    assert "Experience Summary" in extracted
    assert "FastAPI" in extracted
    assert "PostgreSQL" in extracted


def test_extract_text_from_corrupted_docx_rejected():
    corrupted_bytes = b"PK\x03\x04corrupted docx stream not valid zip archive"
    with pytest.raises(ResumeParseError) as exc_info:
        extract_text_from_docx(corrupted_bytes)
    assert exc_info.value.status_code == 400
    assert "unable to read docx" in exc_info.value.message.lower()


def test_extract_text_from_empty_docx_rejected():
    # Document with zero text content
    doc = docx.Document()
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    with pytest.raises(ResumeParseError) as exc_info:
        extract_text_from_docx(docx_bytes)
    assert exc_info.value.status_code == 422
    assert "no extractable text" in exc_info.value.message.lower()


# -----------------------------------------------------------------------------
# High-Level `parse_resume_bytes` Function
# -----------------------------------------------------------------------------

def test_parse_resume_bytes_returns_accurate_word_count():
    sample_text = "One two three four five six seven eight nine ten words."
    pdf_bytes = create_sample_pdf(sample_text)
    text, ftype, word_count = parse_resume_bytes("my_cv.pdf", pdf_bytes)
    assert ftype == "pdf"
    assert word_count >= 10
    assert "words" in text


# -----------------------------------------------------------------------------
# API Endpoint Integration: POST /api/resume/parse
# -----------------------------------------------------------------------------

def test_api_resume_parse_valid_pdf():
    pdf_bytes = create_sample_pdf("Meera Patel Cloud Solutions Architect\nExperienced in AWS and Kubernetes.")
    response = client.post(
        "/api/resume/parse",
        files={"file": ("meera_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "pdf"
    assert "Meera Patel" in data["extracted_text"]
    assert data["word_count"] > 5


def test_api_resume_parse_valid_docx():
    docx_bytes = create_sample_docx(["Sunita Rao Data Engineer", "Expertise in Apache Spark and Python."])
    response = client.post(
        "/api/resume/parse",
        files={"file": ("sunita_resume.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "docx"
    assert "Sunita Rao" in data["extracted_text"]


def test_api_resume_parse_unsupported_format_rejected():
    txt_bytes = b"Just plain text resume"
    response = client.post(
        "/api/resume/parse",
        files={"file": ("resume.txt", io.BytesIO(txt_bytes), "text/plain")},
    )
    assert response.status_code == 400
    assert "unsupported file format" in response.json()["detail"].lower()


def test_api_resume_parse_empty_file_rejected():
    response = client.post(
        "/api/resume/parse",
        files={"file": ("empty.pdf", io.BytesIO(b""), "application/pdf")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_api_resume_parse_corrupt_file_rejected():
    response = client.post(
        "/api/resume/parse",
        files={"file": ("corrupt.pdf", io.BytesIO(b"not a valid pdf content"), "application/pdf")},
    )
    assert response.status_code == 400
    assert "unable to read pdf" in response.json()["detail"].lower()


def test_api_resume_parse_scanned_pdf_rejected_with_422():
    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/api/resume/parse",
        files={"file": ("scanned.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
    )
    assert response.status_code == 422
    assert "scanned image" in response.json()["detail"].lower()
